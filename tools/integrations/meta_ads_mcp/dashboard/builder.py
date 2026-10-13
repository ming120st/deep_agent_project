from __future__ import annotations
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import (
    LineChart,
    Reference,
)
from openpyxl.chart.axis import DateAxis
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from models.meta_ads_analysis import CampaignAnalysis

OUTPUT_DIR = (
    Path(".runtime")
    / "reports"
    / "meta_ads"
)


# =========================================================
# 공통


# =========================================================
# 다양한 타입의 값을 Excel 계산/출력에 사용할 수 있는
# int 또는 float 타입으로 변환한다.


def _to_number(
    value: Any,
) -> int | float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "to_decimal"):
        return float(
            value.to_decimal()
        )
    return float(value)
# 분자/분모가 None이거나 분모가 0이면
# ZeroDivisionError를 발생시키지 않고 None을 반환한다.


def _safe_div(
    numerator: float | None,
    denominator: float | None,
) -> float | None:
    if numerator is None:
        return None
    if denominator is None:
        return None
    if denominator == 0:
        return None
    return numerator / denominator
# 분석 기간 값과 비교 기간 값을 기준으로
# 증감률을 계산한다.
#
# 계산식:
# (분석 기간 값 - 비교 기간 값) / 비교 기간 값


def _calculate_change_rate(
    current: float | None,
    previous: float | None,
) -> float | None:
    if current is None:
        return None
    if previous is None:
        return None
    if previous == 0:
        return None
    return (
        current - previous
    ) / previous
# Excel 표의 Header 셀에 공통 스타일을 적용한다.
#
# - 배경색: 진한 파란색
# - 글자색: 흰색
# - Bold
# - 가로/세로 가운데 정렬


def _style_header(
    cell,
) -> None:
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    cell.font = Font(
        color="FFFFFF",
        bold=True,
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
# Dashboard의 각 섹션 제목에 공통 스타일을 적용한다.
#
# - 연한 파란색 배경
# - Bold
#
# 캠페인 정보, 분석 요약, 성과 분석,
# 액션 영역 등의 제목 스타일에 사용한다.


def _style_section_title(
    cell,
) -> None:
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )
    cell.font = Font(
        bold=True,
    )
# 전달받은 컬럼별 너비 설정값을 기준으로
# Excel Sheet의 컬럼 너비를 일괄 설정한다.


def _set_column_widths(
    sheet,
    widths: dict[str, float],
) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[
            column
        ].width = width


# =========================================================
# 기간 집계


# =========================================================


def _aggregate_snapshots(
    snapshots: list[dict],
) -> dict:
    """
    일별 Snapshot을 기간 단위로 집계한다.
    합산:
    - spend
    - purchases
    - purchase_value
    - impressions
    - clicks
    - link_click
    재계산:
    - CPA
    - ROAS
    - CTR
    - CPC
    - CPM
    - CVR
    - AOV
    주의:
    Reach / Frequency는 일별 값을 합산하여
    정확한 기간 값을 만들 수 없으므로 집계하지 않는다.
    """
    totals = {
        "spend": 0.0,
        "purchases": 0.0,
        "purchase_value": 0.0,
        "impressions": 0.0,
        "clicks": 0.0,
        "link_click": 0.0,
    }
    has_value = {
        key: False
        for key in totals
    }
    for snapshot in snapshots:
        metrics = snapshot.get(
            "metrics",
            {},
        )
        for key in totals:
            value = _to_number(
                metrics.get(key)
            )
            if value is None:
                continue
            totals[key] += value
            has_value[key] = True
    for key in totals:
        if not has_value[key]:
            totals[key] = None
    spend = totals["spend"]
    purchases = totals["purchases"]
    purchase_value = totals[
        "purchase_value"
    ]
    impressions = totals[
        "impressions"
    ]
    clicks = totals["clicks"]
    link_click = totals[
        "link_click"
    ]
    # CPA = 광고비 / 구매
    cost_per_purchase = _safe_div(
        spend,
        purchases,
    )
    # ROAS = 구매매출 / 광고비
    purchase_roas = _safe_div(
        purchase_value,
        spend,
    )
    # CTR = 클릭 / 노출 * 100
    ctr_ratio = _safe_div(
        clicks,
        impressions,
    )
    ctr = (
        ctr_ratio * 100
        if ctr_ratio is not None
        else None
    )
    # CPC = 광고비 / 클릭
    cpc = _safe_div(
        spend,
        clicks,
    )
    # CPM = 광고비 / 노출 * 1000
    cpm_ratio = _safe_div(
        spend,
        impressions,
    )
    cpm = (
        cpm_ratio * 1000
        if cpm_ratio is not None
        else None
    )
    # CVR = 구매 / 링크 클릭 * 100
    cvr_ratio = _safe_div(
        purchases,
        link_click,
    )
    cvr = (
        cvr_ratio * 100
        if cvr_ratio is not None
        else None
    )
    # AOV = 구매매출 / 구매
    aov = _safe_div(
        purchase_value,
        purchases,
    )
    return {
        **totals,
        "cost_per_purchase": (
            cost_per_purchase
        ),
        "purchase_roas": (
            purchase_roas
        ),
        "ctr": ctr,
        "cpc": cpc,
        "cpm": cpm,
        "cvr": cvr,
        "aov": aov,
    }


# =========================================================
# Raw Data Sheet


# =========================================================
RAW_METRIC_COLUMNS = [
    ("광고비", "spend"),
    ("구매", "purchases"),
    ("CPA", "cost_per_purchase"),
    ("구매 전환값", "purchase_value"),
    ("ROAS", "purchase_roas"),
    ("노출", "impressions"),
    ("도달", "reach"),
    ("클릭", "clicks"),
    ("링크 클릭", "link_click"),
    ("CTR", "ctr"),
    ("CPC", "cpc"),
    ("CPM", "cpm"),
    ("빈도", "frequency"),
    ("CVR", "cvr"),
    ("AOV", "aov"),
]


def _create_data_sheet(
    workbook: Workbook,
    current_snapshots: list[dict],
    comparison_snapshots: list[dict],
):
    """
    분석 기간과 비교 기간의 일별 Snapshot을 합친 뒤
    analysis_date 기준 오름차순으로 정렬하여 저장한다.
    기간 구분 없이 가장 오래된 날짜부터 표시한다.
    """
    sheet = workbook.create_sheet(
        "Performance_Data"
    )
    headers = [
        "날짜",
        *[
            label
            for label, _ in RAW_METRIC_COLUMNS
        ],
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        _style_header(cell)
    # 비교 기간 + 분석 기간 데이터를 하나로 합친다.
    all_snapshots = (
        comparison_snapshots
        + current_snapshots
    )
    # 가장 오래된 날짜부터 정렬한다.
    all_snapshots = sorted(
        all_snapshots,
        key=lambda item: item.get(
            "analysis_date",
            "",
        ),
    )
    for snapshot in all_snapshots:
        metrics = snapshot.get(
            "metrics",
            {},
        )
        analysis_date = snapshot.get(
            "analysis_date"
        )
        if analysis_date:
            analysis_date = datetime.strptime(
                analysis_date,
                "%Y-%m-%d",
            )
        sheet.append(
            [
                analysis_date,
                *[
                    _to_number(
                        metrics.get(key)
                    )
                    for _, key in RAW_METRIC_COLUMNS
                ],
            ]
        )
    # 날짜
    for row in range(
        2,
        sheet.max_row + 1,
    ):
        sheet.cell(
            row=row,
            column=1,
        ).number_format = "m/d"
        # 정수/금액 계열
        for col in [
            2, 3, 4, 5,
            7, 8, 9, 10,
            12, 13, 16,
        ]:
            sheet.cell(
                row=row,
                column=col,
            ).number_format = "#,##0"
        # ROAS
        sheet.cell(
            row=row,
            column=6,
        ).number_format = "0.00"
        # CTR
        sheet.cell(
            row=row,
            column=11,
        ).number_format = '0.00"%"'
        # 빈도
        sheet.cell(
            row=row,
            column=14,
        ).number_format = "0.00"
        # CVR
        sheet.cell(
            row=row,
            column=15,
        ).number_format = '0.00"%"'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:P{sheet.max_row}"
    )
    _set_column_widths(
        sheet,
        {
            "A": 14,
            "B": 14,
            "C": 10,
            "D": 14,
            "E": 16,
            "F": 12,
            "G": 12,
            "H": 12,
            "I": 10,
            "J": 12,
            "K": 10,
            "L": 14,
            "M": 14,
            "N": 12,
            "O": 12,
            "P": 14,
        },
    )
    return sheet


# =========================================================
# 날짜 축


# =========================================================


def _configure_date_axis(
    chart: LineChart,
) -> None:
    date_axis = DateAxis(
        axId=10,
        crossAx=100,
    )
    date_axis.axPos = "b"
    date_axis.title = "날짜"
    date_axis.number_format = "m/d"
    date_axis.majorTimeUnit = "days"
    date_axis.tickLblPos = "low"
    date_axis.delete = False
    date_axis.majorUnit = 5
    chart.x_axis = date_axis
    chart.y_axis.axId = 100
    chart.y_axis.crossAx = 10


# =========================================================
# 기간 표시


# =========================================================


def _get_period_text(
    snapshots: list[dict],
) -> str:
    if not snapshots:
        return "-"
    return (
        f"{snapshots[0]['analysis_date']}"
        f" ~ "
        f"{snapshots[-1]['analysis_date']}"
    )


# =========================================================
# Dashboard


# =========================================================


def _create_dashboard_sheet(
    workbook: Workbook,
    current_snapshots: list[dict],
    comparison_snapshots: list[dict],
    current_sheet,
    analysis: CampaignAnalysis,
) -> None:
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    latest = current_snapshots[-1]
    campaign = latest.get(
        "campaign",
        {},
    )
    analysis_data = analysis.model_dump(
        exclude_none=True,
    )
    current_metrics = (
        _aggregate_snapshots(
            current_snapshots
        )
    )
    comparison_metrics = (
        _aggregate_snapshots(
            comparison_snapshots
        )
    )
    current_period = (
        _get_period_text(
            current_snapshots
        )
    )
    comparison_period = (
        _get_period_text(
            comparison_snapshots
        )
    )
    # =====================================================
    # 제목
    # =====================================================
    dashboard.merge_cells(
        "A1:I2"
    )
    title = dashboard["A1"]
    title.value = (
        "Meta Ads Campaign "
        "Period Comparison Dashboard"
    )
    title.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    title.font = Font(
        color="FFFFFF",
        bold=True,
        size=18,
    )
    title.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    # =====================================================
    # 캠페인 / 기간 정보
    # =====================================================
    info_rows = [
        (
            "캠페인",
            campaign.get(
                "campaign_name",
                "",
            ),
        ),
        (
            "분석 기간",
            current_period,
        ),
        (
            "비교 기간",
            comparison_period,
        ),
        (
            "분석 기준일",
            latest.get(
                "analysis_date",
                "",
            ),
        ),
        (
            "목적",
            campaign.get(
                "objective",
                "",
            ),
        ),
        (
            "전환 이벤트",
            campaign.get(
                "conversion_event",
                "",
            ),
        ),
    ]
    for row, (
        label,
        value,
    ) in enumerate(
        info_rows,
        start=4,
    ):
        label_cell = dashboard.cell(
            row=row,
            column=1,
            value=label,
        )
        dashboard.cell(
            row=row,
            column=2,
            value=value,
        )
        _style_section_title(
            label_cell
        )
    # =====================================================
    # 기간 비교 KPI
    # =====================================================
    kpi_start_row = 12
    dashboard.merge_cells(
        start_row=kpi_start_row,
        start_column=1,
        end_row=kpi_start_row,
        end_column=5,
    )
    kpi_title = dashboard.cell(
        row=kpi_start_row,
        column=1,
        value="기간 성과 비교",
    )
    _style_section_title(
        kpi_title
    )
    headers = [
        "지표",
        "분석 기간",
        "비교 기간",
        "증감",
        "증감률",
    ]
    for col, header in enumerate(
        headers,
        start=1,
    ):
        cell = dashboard.cell(
            row=kpi_start_row + 1,
            column=col,
            value=header,
        )
        _style_header(cell)
    compare_items = [
        (
            "광고비",
            "spend",
            "#,##0",
        ),
        (
            "구매",
            "purchases",
            "#,##0",
        ),
        (
            "구매매출",
            "purchase_value",
            "#,##0",
        ),
        (
            "ROAS",
            "purchase_roas",
            "0.00",
        ),
        (
            "CPA",
            "cost_per_purchase",
            "#,##0",
        ),
        (
            "노출",
            "impressions",
            "#,##0",
        ),
        (
            "클릭",
            "clicks",
            "#,##0",
        ),
        (
            "CTR",
            "ctr",
            '0.00"%"',
        ),
        (
            "CPC",
            "cpc",
            "#,##0",
        ),
        (
            "CPM",
            "cpm",
            "#,##0",
        ),
        (
            "CVR",
            "cvr",
            '0.00"%"',
        ),
        (
            "AOV",
            "aov",
            "#,##0",
        ),
    ]
    for index, (
        label,
        key,
        number_format,
    ) in enumerate(
        compare_items,
        start=kpi_start_row + 2,
    ):
        current_value = _to_number(
            current_metrics.get(key)
        )
        comparison_value = _to_number(
            comparison_metrics.get(key)
        )
        difference = None
        if (
            current_value is not None
            and comparison_value is not None
        ):
            difference = (
                current_value
                - comparison_value
            )
        change_rate = (
            _calculate_change_rate(
                current_value,
                comparison_value,
            )
        )
        dashboard.cell(
            row=index,
            column=1,
            value=label,
        )
        current_cell = dashboard.cell(
            row=index,
            column=2,
            value=current_value,
        )
        comparison_cell = dashboard.cell(
            row=index,
            column=3,
            value=comparison_value,
        )
        difference_cell = dashboard.cell(
            row=index,
            column=4,
            value=difference,
        )
        change_cell = dashboard.cell(
            row=index,
            column=5,
            value=change_rate,
        )
        current_cell.number_format = (
            number_format
        )
        comparison_cell.number_format = (
            number_format
        )
        difference_cell.number_format = (
            number_format
        )
        change_cell.number_format = (
            '+0.0%;-0.0%;0.0%'
        )
    # =====================================================
    # 기간 집계 주의사항
    # =====================================================
    note_row = (
        kpi_start_row
        + len(compare_items)
        + 3
    )
    dashboard.merge_cells(
        start_row=note_row,
        start_column=1,
        end_row=note_row,
        end_column=9,
    )
    note_cell = dashboard.cell(
        row=note_row,
        column=1,
        value=(
            "※ CPA, ROAS, CTR, CPC, CPM, CVR, AOV는 "
            "일별 지표의 평균이 아니라 기간 합산 원시값으로 "
            "재계산했습니다. Reach와 Frequency는 일별 값을 "
            "합산하여 정확한 기간 값을 만들 수 없어 "
            "기간 비교 KPI에서 제외했습니다."
        ),
    )
    note_cell.alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )
    note_cell.font = Font(
        italic=True,
        size=9,
    )
    # =====================================================
    # 상세 분석
    # =====================================================
    performance = analysis_data.get(
        "performance",
        {},
    )
    funnel = analysis_data.get(
        "funnel",
        {},
    )
    segment = analysis_data.get(
        "segment",
        {},
    )
    actions = analysis_data.get(
        "actions",
        {},
    )
    limitations = analysis_data.get(
        "limitations",
        [],
    )
    if isinstance(
        limitations,
        list,
    ):
        limitations_text = "\n".join(
            f"• {item}"
            for item in limitations
            if item
        )
    else:
        limitations_text = str(
            limitations or ""
        )
    analysis_sections = [
        (
            "분석 요약",
            analysis_data.get(
                "summary",
                "분석 데이터 없음",
            ),
        ),
        (
            "성과 개요",
            performance.get(
                "overview",
                "분석 데이터 없음",
            ),
        ),
        (
            "효율 분석",
            performance.get(
                "efficiency",
                "분석 데이터 없음",
            ),
        ),
        (
            "볼륨 분석",
            performance.get(
                "volume",
                "분석 데이터 없음",
            ),
        ),
        (
            "노출 · 도달 분석",
            funnel.get(
                "delivery",
                "분석 데이터 없음",
            ),
        ),
        (
            "유입 분석",
            funnel.get(
                "traffic",
                "분석 데이터 없음",
            ),
        ),
        (
            "전환 분석",
            funnel.get(
                "conversion",
                "분석 데이터 없음",
            ),
        ),
        (
            "구매 가치 분석",
            funnel.get(
                "value",
                "분석 데이터 없음",
            ),
        ),
        (
            "광고세트 분석",
            segment.get(
                "adset",
                "분석 데이터 없음",
            ),
        ),
        (
            "소재 분석",
            segment.get(
                "creative",
                "분석 데이터 없음",
            ),
        ),
        (
            "종합 진단",
            analysis_data.get(
                "diagnosis",
                "분석 데이터 없음",
            ),
        ),
        (
            "추가 확인 사항",
            (
                limitations_text
                or "추가 확인 사항 없음"
            ),
        ),
        (
            "캠페인 조치",
            actions.get(
                "campaign",
                "조치 데이터 없음",
            ),
        ),
        (
            "광고세트 조치",
            actions.get(
                "adset",
                "조치 데이터 없음",
            ),
        ),
        (
            "소재 조치",
            actions.get(
                "creative",
                "조치 데이터 없음",
            ),
        ),
        (
            "기술 · 프로덕트 조치",
            actions.get(
                "product",
                "조치 데이터 없음",
            ),
        ),
    ]
    current_row = note_row + 3
    for (
        section_title,
        content,
    ) in analysis_sections:
        dashboard.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=9,
        )
        title_cell = dashboard.cell(
            row=current_row,
            column=1,
            value=section_title,
        )
        _style_section_title(
            title_cell
        )
        content_row = (
            current_row + 1
        )
        dashboard.merge_cells(
            start_row=content_row,
            start_column=1,
            end_row=content_row + 2,
            end_column=9,
        )
        content_cell = dashboard.cell(
            row=content_row,
            column=1,
            value=content,
        )
        content_cell.alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )
        for row in range(
            content_row,
            content_row + 3,
        ):
            dashboard.row_dimensions[
                row
            ].height = 28
        current_row += 5
    # =====================================================
    # 분석 기간 일별 추세 차트
    # =====================================================
    if len(current_snapshots) >= 2:
        comparison_count = len(
            comparison_snapshots
        )
        current_start_row = (
            comparison_count + 2
        )
        current_end_row = (
            current_start_row
            + len(current_snapshots)
            - 1
        )
        # 날짜: A열
        dates = Reference(
            current_sheet,
            min_col=1,
            min_row=current_start_row,
            max_row=current_end_row,
        )
        # ROAS
        roas_chart = LineChart()
        roas_chart.title = (
            "분석 기간 ROAS 일별 추세"
        )
        roas_chart.y_axis.title = "ROAS"
        roas_chart.y_axis.delete = False
        roas_chart.y_axis.tickLblPos = (
            "nextTo"
        )
        roas_chart.y_axis.numFmt = "0.00"
        _configure_date_axis(
            roas_chart
        )
        roas_chart.height = 8
        roas_chart.width = 14
        # ROAS: F열
        roas_data = Reference(
            current_sheet,
            min_col=6,
            min_row=current_start_row,
            max_row=current_end_row,
        )
        roas_chart.add_data(
            roas_data,
            titles_from_data=False,
        )
        roas_chart.set_categories(
            dates
        )
        roas_chart.legend = None
        dashboard.add_chart(
            roas_chart,
            "K4",
        )
        # CPA
        cpa_chart = LineChart()
        cpa_chart.title = (
            "분석 기간 CPA 일별 추세"
        )
        cpa_chart.y_axis.title = "CPA"
        cpa_chart.y_axis.delete = False
        cpa_chart.y_axis.tickLblPos = (
            "nextTo"
        )
        cpa_chart.y_axis.numFmt = "#,##0"
        _configure_date_axis(
            cpa_chart
        )
        cpa_chart.height = 8
        cpa_chart.width = 14
        # CPA: D열
        cpa_data = Reference(
            current_sheet,
            min_col=4,
            min_row=current_start_row,
            max_row=current_end_row,
        )
        cpa_chart.add_data(
            cpa_data,
            titles_from_data=False,
        )
        cpa_chart.set_categories(
            dates
        )
        cpa_chart.legend = None
        dashboard.add_chart(
            cpa_chart,
            "K20",
        )
    # =====================================================
    # Layout
    # =====================================================
    _set_column_widths(
        dashboard,
        {
            "A": 20,
            "B": 24,
            "C": 20,
            "D": 18,
            "E": 18,
            "F": 16,
            "G": 8,
            "H": 16,
            "I": 16,
        },
    )
    dashboard.freeze_panes = "A4"


# =========================================================
# Public Builder


# =========================================================


def build_meta_ads_dashboard(
    current_snapshots: list[dict],
    comparison_snapshots: list[dict],
    analysis: CampaignAnalysis,
) -> str:
    if not current_snapshots:
        raise ValueError(
            "분석 기간 Snapshot이 없습니다."
        )
    if not comparison_snapshots:
        raise ValueError(
            "비교 기간 Snapshot이 없습니다."
        )
    campaign_id = str(
        current_snapshots[-1][
            "campaign"
        ][
            "campaign_id"
        ]
    )

    output_dir = (
        OUTPUT_DIR
        / campaign_id
    )

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )
    current_snapshots = sorted(
        current_snapshots,
        key=lambda item: (
            item.get(
                "analysis_date",
                "",
            )
        ),
    )
    comparison_snapshots = sorted(
        comparison_snapshots,
        key=lambda item: (
            item.get(
                "analysis_date",
                "",
            )
        ),
    )
    workbook = Workbook()
    current_sheet = (
        _create_data_sheet(
            workbook=workbook,
            current_snapshots=current_snapshots,
            comparison_snapshots=(
                comparison_snapshots
            ),
        )
    )
    _create_dashboard_sheet(
        workbook=workbook,
        current_snapshots=(
            current_snapshots
        ),
        comparison_snapshots=(
            comparison_snapshots
        ),
        current_sheet=current_sheet,
        analysis=analysis,
    )
    latest_date = (
        current_snapshots[-1]
        .get(
            "analysis_date",
            "unknown",
        )
    )
    output_file = (
        output_dir
        / (
            "Meta_Ads_Period_Report_"
            f"{campaign_id}_"
            f"{latest_date}.xlsx"
        )
    )
    workbook.save(
        output_file
    )
    return str(
        output_file
    )
